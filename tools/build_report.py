"""Create the faculty report from measured artifacts using python-docx."""
import csv
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'deliverables'
OUT.mkdir(exist_ok=True)
m = json.loads((ROOT/'artifacts/metrics.json').read_text())
s = json.loads((ROOT/'docs/student.json').read_text())
rows = list(csv.DictReader((ROOT/'artifacts/comparison.csv').open()))
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin, sec.bottom_margin = Inches(.75), Inches(.7)
sec.left_margin = sec.right_margin = Inches(.8)
sec.header_distance = sec.footer_distance = Inches(.32)
styles = doc.styles
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    styles[name].font.name = 'Calibri'
    styles[name].font.color.rgb = RGBColor(0,0,0)
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(7)
styles['Normal'].paragraph_format.line_spacing = 1.12
styles['Title'].font.size = Pt(29)
styles['Title'].font.bold = True
styles['Heading 1'].font.size = Pt(21)
styles['Heading 1'].paragraph_format.space_after = Pt(15)
styles['Heading 2'].font.size = Pt(13)
styles['Heading 2'].paragraph_format.space_before = Pt(13)
sec.different_first_page_header_footer = True
header = sec.header.paragraphs[0]
header.text = 'CSE  /  PDEU                                      MACHINE LEARNING LABORATORY'
header.style = 'Caption'
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor(0,0,0)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Payas Vaishnav  •  23BCP430     |     ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
for run in footer.runs: run.font.size = Pt(8)

def p(text, style=None): return doc.add_paragraph(text, style)
def h(text): doc.add_heading(text, 2)
def page(title): doc.add_page_break(); doc.add_heading(title, 1)
def picture(name, width=6.25):
    pp = doc.add_paragraph(); pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pp.add_run().add_picture(str(ROOT/'artifacts'/name), width=Inches(width))
def caption(text):
    pp=p(text,'Caption'); pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
def table(headers, values, widths=None):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for j,v in enumerate(headers): t.rows[0].cells[j].text=str(v)
    for vals in values:
        for cell,v in zip(t.add_row().cells, vals): cell.text=str(v)
    for i,row in enumerate(t.rows):
        for j,cell in enumerate(row.cells):
            if widths: cell.width=Inches(widths[j])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'), '17344D' if i==0 else ('EEF4F7' if i%2 else 'FFFFFF')); pr.append(shade)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
            pr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:w'),'90'); e.set(qn('w:type'),'dxa'); margins.append(e)
            pr.append(margins)
            for pp in cell.paragraphs:
                pp.paragraph_format.space_after=Pt(2); pp.paragraph_format.space_before=Pt(2)
                pp.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in pp.runs:
                    r.font.size=Pt(9); r.font.bold=i==0; r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
        if i==0:
            repeat=OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(repeat)
        no_split=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
    p('')
    return t

p('MACHINE LEARNING LABORATORY', 'Subtitle')
p('Network Intrusion Detection Using Machine Learning','Title')
p('Comparative evaluation of tree classifiers and ensembles on sampled CICIDS2017 network flows','Subtitle')
p('')
p('Semester Project Report', 'Heading 2')
p(s['name'] + '\nEnrollment number ' + s['roll'] + '\n' + s['semester'])
p(s['department']+'\n'+s['college'])
p('Faculty guide\n'+s['guide'])
p('22 October 2026')
h('Project outcome')
p(f"The implemented system compares eight learning configurations, selects {m['best_model']} using validation macro F1, and provides a local dashboard for flow classification. The measured internal test accuracy is {m['test_accuracy']:.2%}, with macro F1 {m['test_macro_f1']:.4f}.")
p('The project includes reproducible training, saved predictions, an executed learning notebook and an interactive demonstration. The evaluation uses a sampled historical benchmark and does not establish live-network reliability.')

page('Abstract and report overview')
p('Network flow records contain measurements that can help distinguish benign communication from malicious traffic. This project develops an offline multiclass intrusion classification system using a sampled CICIDS2017 dataset provided through Western OC2 Lab’s IDS-ML repository. The implementation combines a reproducible experiment pipeline with a local web interface for CSV-based inference.')
p(f"After removing {m['duplicates_removed']:,} duplicate feature rows, the experiment uses {m['rows']:,} flows and {m['features']} measurements. Stratified partitions allocate 60% of the distinct flows to training, 20% to validation and 20% to testing. Median imputation fits inside model pipelines. Experiments cover a majority baseline, Decision Tree, a tuned tree, Random Forest, Extra Trees, Histogram Gradient Boosting, a tree using 20 selected features, and a stacking ensemble.")
p(f"Three-fold training cross-validation supports the comparison, while validation macro F1 determines the selected configuration. The final {m['best_model']} model achieves {m['test_accuracy']:.2%} test accuracy and {m['test_macro_f1']:.4f} macro F1. Per-class scores, binary attack-detection errors and stratified bootstrap intervals provide a more useful assessment than accuracy alone.")
p('The contribution is an integrated educational implementation and comparative evaluation. The work credits its source research and data. The test partition was also evaluated during the initial prototype, so the final results remain an internal benchmark. Future work should evaluate on separate capture days or networks using richer collection metadata.')
h('Report structure')
table(['Section','Focus'],[
('Problem and prior work','Purpose, objectives and relationship to the reference project'),
('Dataset and preparation','Features, class balance and cleaning decisions'),
('Experimental method','Partitions, tuning, model selection and metrics'),
('Models and comparison','Configuration choices and observed tradeoffs'),
('Test evaluation','Confusion matrix, per-class metrics and errors'),
('System and reproduction','Application workflow, rerun commands and limitations')],[2.0,4.5])

page('Problem and prior work')
h('Problem statement')
p('A large collection of network flow records is difficult to inspect manually. The task in this project is to learn a mapping from numeric traffic measurements to one of seven known traffic categories, then apply that mapping consistently to new rows with the same feature schema. The output supports investigation of suspicious traffic rather than automatically blocking communication.')
h('Objectives')
for text in ['Prepare labeled flow data and prevent identical feature rows from crossing partition boundaries.',
             'Compare simple and ensemble classifiers with a majority baseline and class-sensitive metrics.',
             'Study whether hyperparameter tuning and reducing the feature set improve validation performance.',
             'Save the complete fitted preprocessing and classification pipeline for repeatable inference.',
             'Provide a clear demonstration with downloadable predictions and an inspectable error analysis.']:
    p(text,'List Bullet')
h('Relationship to the source research')
p('Western OC2 Lab’s IDS-ML repository contains notebook implementations associated with tree-based IDS, MTH-IDS and LCCDE research. Yang et al. describe a tree-based intrusion detection approach in their GLOBECOM 2019 paper [1]. This project adopts that general application and uses the repository’s sampled CICIDS2017 CSV [2].')
p('The submitted code adds a separate reproducible benchmark and dashboard. It compares scikit-learn models, introduces a baseline, documents partitioning, runs a bounded tree hyperparameter search, and reports errors. It does not claim to reproduce the complete MTH-IDS or LCCDE frameworks, and Histogram Gradient Boosting is not presented as XGBoost.')
h('Scope of the application')
p('Input consists of numeric flow measurements, not packet payloads or raw PCAP files. The application predicts known labels from an uploaded CSV. Live packet capture, traffic blocking, alert delivery and unknown-attack detection are outside the implemented scope. These boundaries keep the demonstration reproducible on an ordinary student laptop.')

page('Dataset and preparation')
p('CICIDS2017 contains benign traffic and multiple attack scenarios with flow measurements derived using CICFlowMeter [3]. This project uses the upstream sample file rather than the complete original dataset. The labels in this sample are BENIGN, Bot, BruteForce, DoS, Infiltration, PortScan and WebAttack.')
table(['Data property','Value'], [('Source sample rows',f"{m['rows']+m['duplicates_removed']:,}"),('Removed duplicate feature rows',f"{m['duplicates_removed']:,}"),('Retained distinct flows',f"{m['rows']:,}"),('Numeric feature columns',m['features']),('Traffic labels',len(m['labels']))],[4.6,1.9])
picture('class_distribution.png',6.2)
caption('Figure 1  Retained class counts on a logarithmic horizontal axis')
p('The rarest label, Infiltration, contains only 36 distinct flows. A perfect result on its small test subset is therefore weak evidence. Macro F1 gives each category equal influence, and class-balanced fitting increases the influence of scarce labels where supported by the classifier.')
p('Preparation trims column names and labels, removes missing labels and exported index columns, converts feature values to numeric form, and treats infinite values as missing. Deduplication keeps the first row for identical feature vectors, including conflicting-label duplicates. Median imputation then fits on each fitting partition. The script retains constant features; univariate feature selection can issue warnings for them, which the saved training log records.')

page('Experimental method')
table(['Partition','Distinct flows','Use'], [('Training',f"{m['split']['train']:,}",'Fit models and perform tree search'),('Validation',f"{m['split']['validation']:,}",'Select configuration by macro F1'),('Test',f"{m['split']['test']:,}",'Report selected model performance')],[1.25,1.35,3.9])
h('Training and selection procedure')
for text in ['Deduplicate flow features and create stratified partitions with random seed 42.',
             'Search 12 tree settings with three-fold stratified cross-validation on the training partition only.',
             'Evaluate all eight configurations with three training folds, then fit each on the complete training partition and predict validation labels.',
             'Select the non-baseline configuration with the highest validation macro F1.',
             'Refit that configuration on training plus validation data, then predict the fixed test partition.',
             'Save metrics, the fitted pipeline, partition identifiers, error rows and example inputs.']:
    p(text,'List Number')
h('Metrics')
p('Accuracy is the number of correctly classified flows divided by the total number of flows. For a class, precision is TP / (TP + FP), recall is TP / (TP + FN), and F1 is 2 × precision × recall / (precision + recall). Macro F1 averages class F1 values equally. A confusion matrix identifies which actual categories the model confuses.')
p('For operational interpretation, labels also map to normal versus attack. The false alarm rate is benign flows predicted as attack divided by all benign flows. Attack recall is detected attack flows divided by all attack flows.')
h('Evaluation qualifications')
p('The fixed sample test split also appears in the initial prototype evaluation. The expanded benchmark keeps selection on validation but does not create an independent external test. Correlated flows may remain across a random split. The tuned tree’s repeated CV score reuses search folds, so it is a diagnostic rather than a nested-CV generalization estimate. Bootstrap intervals reflect resampling of these test examples, not collection-level independence.')

page('Models and configuration choices')
table(['Configuration','Main settings'],[
('Majority baseline','Always predicts the most frequent training label'),
('Decision Tree','Maximum depth 20, balanced class weights'),
('Tuned Decision Tree','12 random candidates across depth, leaf size and split criterion'),
('Random Forest','120 trees, balanced class weights'),
('Extra Trees','120 trees, balanced class weights'),
('Histogram Gradient Boosting','120 maximum iterations, 15 leaves, balanced class weights'),
('Top 20 features Tree','Training-only ANOVA ranking, 20 features, depth 20'),
('Stacking ensemble','Decision Tree and 60-tree Extra Trees; logistic regression meta-model; 3 folds')],[2.2,4.3])
p('The Decision Tree learns feature thresholds. Random Forest averages randomized trees, and Extra Trees adds randomness in threshold construction. Histogram Gradient Boosting builds trees sequentially to improve the current model. Stacking fits a meta-classifier on cross-validated base-model predictions [4]. These methods explore different ways to use nonlinear relationships in tabular measurements.')
h('Hyperparameter search')
p('The tree search considers maximum depth in {8, 12, 20, 30, unlimited}, minimum leaf size in {1, 2, 4, 8}, and criterion in {Gini, entropy}. It samples 12 of the 40 combinations using seed 42 and chooses by training CV macro F1. The selected settings are recorded below.')
for k,v in m['tuning_best_params'].items(): p(k.replace('decisiontreeclassifier__','') + ' = ' + str(v))
p('Feature selection provides a separate ablation: can 20 measurements retain enough information for accurate classification? The selector fits only within the fitting pipeline. It reduces the inputs to the tree, but the dashboard still requires the common full input schema for consistent validation.')

page('Validation and cross validation results')
table(['Configuration','Val accuracy','Val macro F1','CV mean ± SD'],[
    (r['model'],f"{float(r['validation_accuracy']):.2%}",f"{float(r['validation_macro_f1']):.4f}",f"{float(r['cv_macro_f1_mean']):.4f} ± {float(r['cv_macro_f1_std']):.4f}") for r in rows],[2.35,1.1,1.1,1.95])
picture('model_comparison.png',6.25)
caption('Figure 2  Validation macro F1 for all eight configurations')
p('Histogram Gradient Boosting narrowly exceeds the original Decision Tree on validation macro F1. Stacking and the 20-feature tree score lower in this experiment. The majority baseline achieves about 49% validation accuracy while its macro F1 is only about 0.0943, showing why accuracy alone is insufficient. The observed CV variation also cautions against treating small score differences as conclusive.')

page('Test performance and confusion matrix')
table(['Selected model','Test accuracy','Test macro F1'],[(m['best_model'],f"{m['test_accuracy']:.2%}",f"{m['test_macro_f1']:.4f}")],[3.3,1.6,1.6])
picture('confusion_matrix.png',5.65)
caption('Figure 3  Test confusion matrix with actual labels on rows and predictions on columns')
p(f"The selected model makes {m['split']['test']-sum(m['confusion_matrix'][i][i] for i in range(len(m['labels'])))} multiclass mistakes among {m['split']['test']:,} test examples. Most observations lie on the diagonal. Off-diagonal cells show both benign-versus-attack errors and confusions between attack categories.")
p(f"A 300-resample stratified bootstrap gives an accuracy interval of {m['accuracy_ci95'][0]:.2%} to {m['accuracy_ci95'][1]:.2%} and a macro F1 interval of {m['macro_f1_ci95'][0]:.4f} to {m['macro_f1_ci95'][1]:.4f}. These are descriptive 95% percentile intervals for this sample, without guarantees under dataset shift.")

page('Per class results and missed attacks')
table(['Traffic label','Precision','Recall','F1','Support'],[(label,f"{m['classification_report'][label]['precision']:.4f}",f"{m['classification_report'][label]['recall']:.4f}",f"{m['classification_report'][label]['f1-score']:.4f}",int(m['classification_report'][label]['support'])) for label in m['labels']],[1.65,1.2,1.2,1.2,1.25])
b=m['binary_confusion_matrix']
table(['Actual traffic','Predicted normal','Predicted attack'],[('Normal',b[0][0],b[0][1]),('Attack',b[1][0],b[1][1])],[2.5,2.0,2.0])
p(f"The model flags {b[0][1]} benign flows as attacks and misses {b[1][0]} attack flows by predicting BENIGN. Attack recall is {m['binary_attack_recall']:.2%}; the benign false alarm rate is {m['binary_false_positive_rate']:.2%}. Other multiclass errors confuse attack categories without missing their attack status.")
p('A false alarm can increase investigation workload. A missed attack can allow malicious activity to pass unnoticed. This implementation uses the classifier’s standard predicted label; it does not optimize a cost-sensitive alert threshold. Its confidence values are not calibrated risk probabilities.')
weak=min(m['labels'], key=lambda k:m['classification_report'][k]['f1-score'])
p(f"{weak} has the lowest per-class test F1 in this run. Infiltration has only seven test examples, so its score cannot establish reliable detection of that class. The CSV in artifacts/error_examples.csv retains feature measurements with actual and predicted labels for inspection.")
h('Practical interpretation')
p('The measured scores support the feasibility of classification within this sampled benchmark. They do not prove the system will detect a novel attack or remain accurate when traffic characteristics change. Independent evaluation by capture day, device or network is the most useful next validation step.')

page('Feature analysis and application design')
picture('feature_importance.png',6.1)
caption('Figure 4  Diagnostic Decision Tree feature importance for interpretation')
p('Histogram Gradient Boosting does not expose impurity-based feature importance through the interface used here. The chart therefore comes from a separate diagnostic Decision Tree fitted on the development partition. It is explicitly labeled as a diagnostic model. It does not explain the selected boosting model or establish causal importance.')
h('Application components')
table(['Module','Responsibility'],[('ids.py','Validate CSVs, clean numeric values and run saved-model inference'),('train.py','Tune, compare, select, evaluate and save artifacts'),('app.py','Display results, experiments and downloadable CSV predictions'),('verify.py','Check inference, schema rejection and dashboard demo interaction'),('IDS_Experiment.ipynb','Executed walkthrough of dataset, results and predictions')],[2.15,4.35])
p('The dashboard provides Model results, Detect attacks, Experiments and Understand the project tabs. Inference loads the locally generated pipeline, requires matching feature columns, limits uploads to 50,000 rows, and exports predictions with labels and confidence. Uploaded files are processed in memory; the application does not add them to the training set.')

page('Reproduction and demonstration')
h('Reproduce the experiment')
for text in ['Install the pinned packages from requirements-lock.txt in a Python 3.13 virtual environment. The verified environment uses scikit-learn '+m['sklearn']+'.',
             'Ensure reference/data/CICIDS2017_sample.csv is present. setup.ps1 can download the upstream repository.',
             'Run .venv\\Scripts\\python.exe train.py, then .venv\\Scripts\\python.exe verify.py.',
             'Run .venv\\Scripts\\python.exe prepare_materials.py to refresh the measured results guide and executed notebook.',
             'Start the dashboard with .venv\\Scripts\\python.exe -m streamlit run app.py. Restart after retraining to clear the model cache.']:
    p(text,'List Number')
h('Evidence retained')
p('metrics.json records final scores, bootstrap intervals, versions, seed and data checksum. comparison.csv stores validation and CV results. tuning.csv records all 12 search candidates. partitions.csv provides deterministic row-to-partition assignments. test_predictions.csv and error_examples.csv allow independent metric and error checks. The training log records execution output and feature-selection warnings.')
p('Dataset SHA-256', 'Caption')
pp=p(m['dataset_sha256']); pp.runs[0].font.size=Pt(8)
h('Demonstration sequence')
for text in ['Explain the traffic classification problem and the seven labels.',
             'Show the baseline comparison, validation selection and test confusion matrix.',
             'Open Detect attacks, retain demo flows and click Analyze flows. These are held-out test rows.',
             'Download the predictions, then show error groups in the Experiments tab.',
             'Explain one limitation and the next independent evaluation you would perform.']:
    p(text,'List Bullet')
p('The supplied demo flows are held-out examples with target labels removed from the input. After setup, local model inference works without internet access. Keep the trained artifacts and pinned environment available on the demonstration computer.')

page('Conclusions and references')
h('Conclusions')
p(f"The project delivers a working offline intrusion classification system with a reproducible benchmark and a usable interface. Among eight evaluated configurations, {m['best_model']} achieves the highest validation macro F1. The selected model reaches {m['test_accuracy']:.2%} accuracy and {m['test_macro_f1']:.4f} macro F1 on the fixed internal test split.")
p('The experiments also reveal useful negative results: tuning does not automatically improve validation performance, a small feature subset loses information, and the tested stacking configuration underperforms the simpler tree. Those findings support choosing a model from evidence rather than complexity alone.')
h('Limitations and future work')
p('The sample is historical, rare classes are small, exact deduplication cannot remove all correlated observations, and the test split has been used previously. Bootstrap intervals do not model these effects. Future work should use full capture metadata to evaluate by day or session, test a separate deployment network, calibrate confidence values on separate data, and measure live extraction and inference latency before integrating alerts.')
h('References')
for text in [
    '[1] L. Yang, A. Moubayed, I. Hamieh and A. Shami. Tree-Based Intelligent Intrusion Detection System in Internet of Vehicles. IEEE GLOBECOM, 2019. DOI 10.1109/GLOBECOM38437.2019.9013892.',
    '[2] Western OC2 Lab. Intrusion Detection System Using Machine Learning. https://github.com/Western-OC2-Lab/Intrusion-Detection-System-Using-Machine-Learning. Sample and reference revision 502ba86173a8aae5a9b2973fa4e0a42de69fa6ba. Upstream MIT license retained in reference/LICENSE.',
    '[3] Canadian Institute for Cybersecurity, University of New Brunswick. CICIDS2017 dataset documentation. https://www.unb.ca/cic/datasets/ids-2017.html. Accessed 22 October 2026.',
    '[4] scikit-learn documentation. Ensemble methods. https://scikit-learn.org/stable/modules/ensemble.html. Accessed 22 October 2026.',
    '[5] scikit-learn documentation. Cross-validation. https://scikit-learn.org/stable/modules/cross_validation.html. Accessed 22 October 2026.'
]:
    pp=p(text); pp.paragraph_format.space_after=Pt(10); pp.runs[0].font.size=Pt(9)
doc.core_properties.title='Network Intrusion Detection Using Machine Learning'
doc.core_properties.author=s['name']
doc.core_properties.subject='ML laboratory semester project report'
doc.save(OUT/'IDS_Project_Report_Editable.docx')
print('Report created:',OUT/'IDS_Project_Report_Editable.docx')
